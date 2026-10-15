""" Module for checking file and datasheet upload"""
from werkzeug.utils import secure_filename
import os
import pandas as pd
from collections import Counter, defaultdict, OrderedDict
from dataclasses import dataclass, asdict
from html import escape
import datetime
import logging
import math
import ntpath
import json
import re
import openpyxl
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# NB the datasheet that passes these checks is NOT what the framework loads. The original
# uploaded file is moved as is to the submission directory and the framework reads it again
# with its own code (symportal_framework/data_loading.py, _read_in_datasheet).
# Therefore everything here must read the file in the same way as the framework does, and
# may only be stricter than the framework in what it accepts.

EXPECTED_FIRST_SHEET = 'REAL_DATA'
EXPECTED_OTHER_SHEETS = ['EXAMPLE_DATA', 'IGNORE']

EXPECTED_COLUMNS = [
    'sample_name', 'fastq_fwd_file_name', 'fastq_rev_file_name',
    'sample_type', 'host_phylum', 'host_class', 'host_order',
    'host_family', 'host_genus', 'host_species',
    'collection_latitude', 'collection_longitude',
    'collection_date', 'collection_depth'
]

# The accepted sample types. These are the values listed in the IGNORE tab of the template.
SAMPLE_TYPES = [
    'coral_field', 'coral_aquarium', 'aiptasia', 'other_anemone', 'clam', 'sponge',
    'other_animal_host', 'culture', 'seawater', 'sediment', 'epiphytic', 'PCR_control',
    'other'
]

HOST_FIELDS = ['host_phylum', 'host_class', 'host_order', 'host_family', 'host_genus', 'host_species']
LAT = 'collection_latitude'
LON = 'collection_longitude'
# Every column after sample_type. sample_type itself is always required.
METADATA_FIELDS = HOST_FIELDS + [LAT, LON, 'collection_date', 'collection_depth']

# The fields that do not exist for a given sample type and may therefore be left blank.
# Any sample type that is not listed here has every field required.
# TODO define for aiptasia, seawater, sediment and epiphytic.
FIELDS_NOT_APPLICABLE = {
    'PCR_control': HOST_FIELDS + [LAT, LON, 'collection_date', 'collection_depth'],
    'culture': HOST_FIELDS,
    'coral_aquarium': [LAT, LON, 'collection_depth'],
}

# The only text values that we accept in place of a blank cell, and only for a field that
# does not exist for the sample type. The framework converts exactly these to null.
FRAMEWORK_NULL_TOKENS = {'N/A', 'NA', 'na', 'n/a'}

# The strings that pandas reads as null by default. The framework reads the datasheet with
# these defaults so a cell holding one of these is empty as far as the framework is concerned.
# We read with keep_default_na=False so that we see what the user actually typed.
PANDAS_NA_STRINGS = {
    '-1.#IND', '1.#QNAN', '1.#IND', '-1.#QNAN', '#N/A N/A', '#N/A', 'N/A', 'n/a', 'NA',
    '<NA>', '#NA', 'NULL', 'null', 'NaN', '-NaN', 'nan', '-nan', 'None'
}

# Anything that stands in for an empty cell. Compared after strip and lower.
EMPTY_LOOKING = {_.lower() for _ in PANDAS_NA_STRINGS} | {'na', 'none', '-', '--', '?', 'unknown', 'nodata'}

# Lengths of the corresponding fields of the DataSetSample model of the framework.
MAX_FIELD_LENGTHS = {
    'sample_name': 200, 'fastq_fwd_file_name': 100, 'fastq_rev_file_name': 100,
    'sample_type': 50, 'host_phylum': 50, 'host_class': 50, 'host_order': 50,
    'host_family': 50, 'host_genus': 50, 'host_species': 50,
    'collection_date': 40, 'collection_depth': 40
}

VALID_SAMPLE_NAME_CHARS = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_.")

# How a date cell of excel looks once it has been read in as text
EXCEL_DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}( \d{2}:\d{2}:\d{2}(\.\d+)?)?$')
# e.g. 24.339216 or -53.059163. The template asks for decimal degrees. The framework can convert some other
# formats, e.g. 24.33N or 48°53'10.18"N, but we do not accept these.
DECIMAL_DEGREES_RE = re.compile(r'^[+-]?(\d+(\.\d*)?|\.\d+)$')
DEPTH_RE = re.compile(r'^\d+(\.\d+)?(\s*-\s*\d+(\.\d+)?)?$')

# The four things that a cell can hold
VALUE = 'value'
BLANK = 'blank'
FRAMEWORK_NULL = 'framework_null'
STAND_IN = 'stand_in'

# The headings under which the problems are reported to the user
MSG_REQUIRED_EMPTY = 'Required cell is empty'
MSG_LEAVE_BLANK = 'Leave the cell blank'
MSG_NAME_MISSING = 'Sample name missing'
MSG_NAME_CHARS = 'Sample name: only letters, numbers, - _ . allowed'
MSG_NAME_DUPLICATE = 'Sample name used more than once'
MSG_NAME_CHANGED = 'Sample names will be changed'
MSG_FWD_MISSING = 'Forward sequence file name missing'
MSG_REV_MISSING = 'Reverse sequence file name missing'
MSG_FILE_PATH = 'Give the file name only, not the full path'
MSG_FILE_ENDING = 'File name must end in .fastq.gz or .fq.gz'
MSG_FILE_DUPLICATE = 'Sequence file name used more than once'
MSG_TYPE_NOT_ACCEPTED = 'sample_type not accepted (see the list in the template)'
MSG_TYPE_CASE = 'sample_type capitalisation differs from the template'
MSG_DATE_FORMAT = 'Date must be YYYYMMDD, e.g. 20200618'
MSG_DATE_NOT_VALID = 'Date not valid: check the year, month and day'
MSG_DATE_EXCEL = 'Cell is an Excel date: format it as Text and enter YYYYMMDD, e.g. 20200618'
MSG_DEPTH_FORMAT = 'Depth must be a number or a range, e.g. 6 or 5-12'
MSG_DEPTH_EXCEL = 'Cell is an Excel date: format it as Text and enter the depth, e.g. 5-12'
MSG_LAT_RANGE = 'Latitude out of range (-90 to 90)'
MSG_LON_RANGE = 'Longitude out of range (-180 to 180)'
MSG_LAT_LON_NOT_DECIMAL = 'Latitude/longitude must be decimal degrees, e.g. 24.339216'
MSG_LAT_LON_BOTH = 'Latitude and longitude must both be given'
MSG_SPECIES_SHORTENED = 'host_species will be shortened to the species name'
MSG_NOT_XLSX = 'The datasheet must be an Excel .xlsx file'
MSG_CANNOT_OPEN = 'The datasheet could not be opened. Please use the .xlsx template'
MSG_NO_HEADINGS = 'Column headings not found on row 2 of the first tab. Please use the template'
MSG_HEADING_MISSING = 'Column heading missing'
MSG_HEADING_DUPLICATE = 'Column heading used more than once'
MSG_HEADING_OUTSIDE = 'Column must be within columns A to N'
MSG_COLUMN_ORDER = 'Columns are not in the template order'
MSG_EXTRA_COLUMNS = 'Columns after N are not imported'
MSG_EXTRA_TABS = 'Extra tabs are ignored'
MSG_NO_SAMPLES = f'No samples found. Enter them in the first tab ({EXPECTED_FIRST_SHEET}) from row 3'


# TODO QC
# Staged same file twice

# Most modern way of defining advanced exceptions here:
# https://stackoverflow.com/a/53469898/5516420
class DatasheetGeneralFormattingError(Exception):
    """Exception raised when a general formatting error has occured
        For example, if non-unique sample names are used
    """
    def __init__(self, message, data=None):
        self.message = message
        self.data = data

    def __str__(self):
        return str(self.message)

class AddedFilesError(Exception):
    """
    Raised when there are either files missing,
    there are files that are not in the supplied datasheet
    or if there are files that are too small.
    """
    def __init__(self, message, data):
        self.message = message
        self.data = data

    def __str__(self):
        return str(self.message)

class DateFormatError(Exception):
    """
    Raised when there is an error in the date format in the
    supplied datasheet.
    """
    def __init__(self, message, data):
        self.message = message
        self.data = data

    def __str__(self):
        return str(self.message)

class LatLonError(Exception):
    """
    Raised when there is an error in the latitude or longitude provided
    in the supplied datasheet.
    """
    def __init__(self, message, data):
        self.message = message
        self.data = data

    def __str__(self):
        return str(self.message)

class UploadedFilesError(Exception):
    """
    Raised when something has gone wrong when uploading seq files to the server
    """
    def __init__(self, message, data):
        self.message = message
        self.data = data

    def __str__(self):
        return str(self.message)


@dataclass
class Issue:
    """A single problem found in the datasheet."""
    severity: str
    # The heading that the issue is reported under. Issues of the same kind are reported together.
    kind: str
    # The excel cell reference e.g. K7
    cell: str = None
    # What the user entered
    value: str = None
    # What the value will be changed to
    new_value: str = None
    # e.g. the sample type of the row
    note: str = None
    # The heading of the column of the cell. Reported where the kind does not already say which field it is.
    column: str = None


class IssueCollector:
    """Collects every problem found in the datasheet so that they can be reported in one go,
    grouped by the type of problem."""
    MAX_VALUE_LENGTH = 40

    def __init__(self):
        self.errors = []
        self.warnings = []

    def error(self, kind, cell=None, value=None, new_value=None, note=None, column=None):
        self.errors.append(Issue('error', kind, cell, value, new_value, note, column))

    def warning(self, kind, cell=None, value=None, new_value=None, note=None, column=None):
        self.warnings.append(Issue('warning', kind, cell, value, new_value, note, column))

    def as_dicts(self):
        return [asdict(_) for _ in self.errors + self.warnings]

    def errors_html(self):
        """The message shown to the user when the datasheet has problems that must be fixed."""
        count = len(self.errors)
        if count == 1:
            html = '<strong class="text-danger">1 problem found. Fix it and upload again.</strong><br><br>'
        else:
            html = f'<strong class="text-danger">{count} problems found. Fix them and upload again.</strong><br><br>'
        html += self._issues_to_html(self.errors)
        if self.warnings:
            html += '<br><strong>Warnings</strong><br><br>' + self._issues_to_html(self.warnings)
        return html

    def warnings_html(self, heading=True):
        if not self.warnings:
            return ''
        html = ''
        if heading:
            html += '<strong>Warnings (you can still submit)</strong><br><br>'
        return html + self._issues_to_html(self.warnings)

    def _issues_to_html(self, issues):
        by_kind = OrderedDict()
        for issue in issues:
            by_kind.setdefault(issue.kind, []).append(issue)
        html = ''
        for kind, kind_issues in by_kind.items():
            # We collect the issues by note so that e.g. the cells of one sample type are listed together
            by_note = OrderedDict()
            for issue in kind_issues:
                if issue.cell is None and issue.value is None:
                    continue
                by_note.setdefault(issue.note, []).append(issue)
            groups = []
            for note, note_issues in by_note.items():
                group = ', '.join(self._entries(note_issues))
                columns = list(OrderedDict.fromkeys(_.column for _ in note_issues if _.column))
                if columns:
                    group = f'{escape(", ".join(columns), quote=False)}: {group}'
                if note:
                    group += f' ({escape(note, quote=False)})'
                groups.append(group)
            html += escape(kind, quote=False)
            if groups:
                html += f'<pre>    {"   ".join(groups)}</pre>'
            else:
                html += '<br>'
        return html

    def _entries(self, issues):
        if all(_.value is None for _ in issues):
            return [escape(_, quote=False) for _ in self._compress_cells([_.cell for _ in issues])]
        entries = []
        for issue in issues:
            parts = []
            if issue.cell is not None:
                parts.append(escape(issue.cell, quote=False))
            if issue.value is not None:
                parts.append(f'"{escape(self._shorten(issue.value), quote=False)}"')
            if issue.new_value is not None:
                parts.append(f'-&gt; "{escape(self._shorten(issue.new_value), quote=False)}"')
            entries.append(' '.join(parts))
        return entries

    def _shorten(self, value):
        value = str(value)
        if len(value) > self.MAX_VALUE_LENGTH:
            return value[:self.MAX_VALUE_LENGTH] + '...'
        return value

    @staticmethod
    def _compress_cells(cells):
        """Report runs of cells in the same column as a range, i.e. K7, K8, K9 as K7-K9,
        so that a long datasheet with the same problem in every row stays readable."""
        rows_by_column = OrderedDict()
        other = []
        for cell in cells:
            match = re.match(r'^([A-Z]+)(\d+)$', cell)
            if match:
                rows_by_column.setdefault(match.group(1), []).append(int(match.group(2)))
            else:
                other.append(cell)
        compressed = []
        for column, rows in rows_by_column.items():
            rows = sorted(set(rows))
            start = previous = rows[0]
            for row in rows[1:] + [None]:
                if row is not None and row == previous + 1:
                    previous = row
                    continue
                if start == previous:
                    compressed.append(f'{column}{start}')
                else:
                    compressed.append(f'{column}{start}-{column}{previous}')
                start = previous = row
        return compressed + other


def classify_cell(raw):
    """This is the only place where we decide whether a cell holds a value or is empty.
    Returns one of
        VALUE - a real entry
        BLANK - nothing in the cell
        FRAMEWORK_NULL - one of the exact spellings that the framework converts to null
        STAND_IN - any other text that stands in for an empty cell e.g. 'none', '-', ' N/a', '  '
    """
    if raw is None:
        return BLANK
    if isinstance(raw, float) and math.isnan(raw):
        return BLANK
    raw = str(raw)
    if raw == '':
        return BLANK
    if raw in FRAMEWORK_NULL_TOKENS:
        return FRAMEWORK_NULL
    if raw.strip() == '' or raw.strip().lower() in EMPTY_LOOKING:
        return STAND_IN
    return VALUE


def accepted_date_parts(date_value):
    """The year, month and day of a collection_date (month and day are None if they are not given).
    Returns None if the value is not in one of the formats that we accept:
    YYYYMMDD, YYYYMM, YYYY, or DD.MM.YYYY, which the framework converts to YYYYMMDD.
    This is stricter than the framework, which also accepts e.g. June 2020, DD.MM.YY or 2020061812."""
    date_value = str(date_value).strip()
    # A number cell of excel may come in as e.g. 20200618.0. The framework removes the .0
    if re.fullmatch("[0-9]+\\.0+", date_value):
        date_value = date_value.split(".")[0]
    if re.fullmatch("[0-9]{2}\\.[0-9]{2}\\.[0-9]{4}", date_value):
        day, month, year = date_value.split(".")
        return int(year), int(month), int(day)
    if re.fullmatch("[0-9]{8}", date_value):
        return int(date_value[:4]), int(date_value[4:6]), int(date_value[6:])
    if re.fullmatch("[0-9]{6}", date_value):
        return int(date_value[:4]), int(date_value[4:]), None
    if re.fullmatch("[0-9]{4}", date_value):
        return int(date_value), None, None
    return None


def date_is_real(year, month, day):
    """True if the date exists, is from 1900 onwards and is not in the future."""
    try:
        date = datetime.date(year, month or 1, day or 1)
    except ValueError:
        return False
    return datetime.date(1900, 1, 1) <= date <= datetime.date.today()


def read_datasheet(datasheet_path):
    """Read the datasheet in exactly the same way that the framework does:
    the first tab, the column headings on row 2, columns A to N and every cell as text.
    The only difference is that we do not let pandas convert e.g. 'NA' to null (keep_default_na=False)
    so that we can see what the user actually typed. An empty cell is therefore ''.
    The index of a row + 3 is its row number in excel.
    Rows that are completely empty are dropped.
    """
    if not datasheet_path.endswith('.xlsx'):
        raise RuntimeError(f'Data sheet: {datasheet_path} is in an unrecognised format. '
                           f'Please ensure that it is in .xlsx format.')
    df = pd.read_excel(
        io=datasheet_path, header=0, usecols='A:N', skiprows=[0], dtype=str, keep_default_na=False)
    is_empty_row = df.apply(lambda row: all(classify_cell(_) == BLANK for _ in row), axis=1)
    if len(df.index):
        df = df[~is_empty_row]
    return df


class DatasheetChecker:
    """
    This class is responsible for checking the format of the datasheet and also for checking the contents
    Every problem that is found in the datasheet is collected (self.issues) and they are reported together.
    Errors must be fixed by the user. Warnings do not prevent the user from submitting.
    Out of this class we also want to out put:
    extra objects - These are seq files that have been uploaded but are not included in the datasheet
    missing objects - These are the seq files that are missing.
    This should be output as a json.
    """

    def __init__(self, request, user_upload_directory, datasheet_path=None):
        self.request = request
        self.datasheet_path = datasheet_path
        self.user_upload_directory = user_upload_directory
        os.makedirs(self.user_upload_directory, exist_ok=True)
        self.duplication_dict = {}
        if not self.datasheet_path:
            # Then we need to get the datasheet from the uploaded files
            # and save it to disk
            self.datasheet_filename = self._get_and_save_datasheet()

        # This is populated when the datasheet is checked.
        # It is indexed by sample name.
        self.sample_meta_info_df = None
        self.issues = IssueCollector()
        # The list of the files that are in the datasheet but not found in the local dir
        # Key will be sample name value will be list of the files that are missing
        self.missing_files = defaultdict(list)
        # The list of the files that have been uploaded but that have not been
        self.extra_files = []
        # We must ensure that the same file has not been staged more than once
        # Key is file name, value is number of times it appears in the staged area
        self.duplicate_staged_files = defaultdict(int)
        # Warning containers
        # An empty required cell is now an error and a cell that is empty because the field does not
        # exist for the sample type is not reported, so these are no longer populated.
        # They are kept so that the response sent to the browser keeps its structure.
        self.sample_type_missing_list = []
        self.taxonomy_missing_set = set()
        self.depth_missing_list = []
        self.date_missing_list = []
        self.lat_lon_missing = []
        self.binomial_dict = {}
        self.lat_long_dict = {}
        self.date_dict = {}

    def do_general_format_check(self):
        """Run when the datasheet is uploaded.
        Raises a DatasheetGeneralFormattingError that lists every problem in the datasheet.
        It will be handled in routes.py"""
        self._check_datasheet()

    def check_valid_seq_files_added(self):
        """Run every time the staged seq files change."""
        self._check_datasheet()

        # If this check fails, an AddedFilesError is raised and this will cause an error message
        # to be reported to the user. The user will have to fix the errors before they can upload
        # Errors will be caused by missing seq files, extra seq files, seq files that are too small
        self._check_seq_files_exist()

    def _check_datasheet(self):
        self.issues = IssueCollector()

        # First the things that have to be right for us to be able to look at the samples at all
        self._check_structure()
        self._raise_if_errors()

        df = read_datasheet(self.datasheet_path)
        if df.empty:
            self.issues.error(MSG_NO_SAMPLES)
            self._raise_if_errors()

        # Then every sample
        sample_names = self._check_rows(df)
        self._raise_if_errors()

        df = df.copy()
        df['sample_name'] = sample_names
        for col in ['fastq_fwd_file_name', 'fastq_rev_file_name']:
            df[col] = df[col].str.strip()
        for col in ['sample_type'] + METADATA_FIELDS:
            df[col] = [_.strip() if classify_cell(_) == VALUE else 'NoData' for _ in df[col]]
        self.sample_meta_info_df = df.set_index('sample_name', drop=True)

    def _raise_if_errors(self):
        if self.issues.errors:
            for issue in self.issues.errors:
                logger.info(f'datasheet error: {issue}')
            raise DatasheetGeneralFormattingError(
                message=self.issues.errors_html(),
                data={"error_type": "datasheet_issues", "issues": self.issues.as_dicts()}
            )

    def _check_structure(self):
        """Check the file against the template.
        NB This is only used to report problems. The samples are always read from the first tab with the
        headings on row 2, as the framework does, irrespective of what is found here."""
        if not self.datasheet_path.endswith('.xlsx'):
            self.issues.error(MSG_NOT_XLSX)
            return
        try:
            workbook = openpyxl.load_workbook(self.datasheet_path, data_only=True)
            first_sheet = workbook.worksheets[0]
        except Exception:
            logger.exception(f'unable to open {self.datasheet_path}')
            self.issues.error(MSG_CANNOT_OPEN)
            return

        self._check_tabs(workbook, first_sheet)

        rows = list(first_sheet.iter_rows(values_only=True))
        headings = [None if _ is None else str(_) for _ in (rows[1] if len(rows) > 1 else [])]
        headings_a_to_n = (headings + [None] * len(EXPECTED_COLUMNS))[:len(EXPECTED_COLUMNS)]
        headings_after_n = headings[len(EXPECTED_COLUMNS):]

        self._check_headings(rows, headings_a_to_n, headings_after_n)

        # Columns to the right of N are allowed but they are not used
        extra_columns = []
        for row in rows[1:]:
            for i, cell in enumerate(row[len(EXPECTED_COLUMNS):]):
                column = get_column_letter(len(EXPECTED_COLUMNS) + i + 1)
                if classify_cell(cell) != BLANK and column not in extra_columns:
                    extra_columns.append(column)
        for column in sorted(extra_columns, key=lambda _: (len(_), _)):
            self.issues.warning(MSG_EXTRA_COLUMNS, cell=column)

    def _check_tabs(self, workbook, first_sheet):
        if first_sheet.title != EXPECTED_FIRST_SHEET:
            self.issues.warning(
                f'First tab is "{first_sheet.title}", not "{EXPECTED_FIRST_SHEET}". '
                f'Please check your sample data is in "{first_sheet.title}" or this could affect your submission.'
            )
        tab_names = [_.title for _ in workbook.worksheets]
        for tab_name in EXPECTED_OTHER_SHEETS:
            if tab_name not in tab_names:
                self.issues.warning(f'Tab "{tab_name}" is missing')
        for tab_name in tab_names[1:]:
            if tab_name not in EXPECTED_OTHER_SHEETS + [EXPECTED_FIRST_SHEET]:
                self.issues.warning(MSG_EXTRA_TABS, value=tab_name)

    def _check_headings(self, rows, headings_a_to_n, headings_after_n):
        if headings_a_to_n == EXPECTED_COLUMNS:
            return
        if 'sample_name' not in headings_a_to_n:
            # Then the headings are not where they should be. See if they are on another row.
            for i, row in enumerate(rows[:20]):
                if 'sample_name' in [str(_) for _ in row if _ is not None]:
                    self.issues.error(f'Column headings are on row {i + 1}. They must be on row 2')
                    return
            self.issues.error(MSG_NO_HEADINGS)
            return
        for heading, count in Counter(headings_a_to_n).items():
            if heading in EXPECTED_COLUMNS and count > 1:
                self.issues.error(MSG_HEADING_DUPLICATE, value=heading)
        missing = [_ for _ in EXPECTED_COLUMNS if _ not in headings_a_to_n]
        for heading in missing:
            i = EXPECTED_COLUMNS.index(heading)
            found = headings_a_to_n[i]
            if heading in headings_after_n:
                self.issues.error(MSG_HEADING_OUTSIDE, value=heading)
            elif found is not None and found not in EXPECTED_COLUMNS:
                # Then a different heading is in the place of the one that we expect
                self.issues.error(f'Column {get_column_letter(i + 1)} heading is "{found}", should be "{heading}"')
            else:
                self.issues.error(MSG_HEADING_MISSING, value=heading)
        if not self.issues.errors:
            # All of the headings are there but in a different order. The framework works with
            # the headings so this is OK.
            self.issues.warning(MSG_COLUMN_ORDER)

    def _check_rows(self, df):
        """Check every sample of the datasheet. Returns the list of sample names as they will be used."""
        column_letters = {col: get_column_letter(i + 1) for i, col in enumerate(df.columns)}
        sample_names = []
        # key is the name, value is list of the cells that it is found in
        cells_of_sample_name = defaultdict(list)
        cells_of_file_name = defaultdict(list)

        for df_ind, row in df.iterrows():
            excel_row = df_ind + 3

            def cell(col):
                return f'{column_letters[col]}{excel_row}'

            # Sample name
            sample_name = self._check_sample_name(row['sample_name'], cell('sample_name'))
            sample_names.append(sample_name)
            if sample_name is not None:
                cells_of_sample_name[sample_name].append(cell('sample_name'))

            # Seq file names
            for col, missing_message in [
                ('fastq_fwd_file_name', MSG_FWD_MISSING), ('fastq_rev_file_name', MSG_REV_MISSING)
            ]:
                file_name = self._check_file_name(row[col], cell(col), missing_message)
                if file_name is not None:
                    cells_of_file_name[file_name].append(cell(col))

            # Sample type
            sample_type = self._check_sample_type(row['sample_type'], cell('sample_type'))

            # Empty cells
            # If we were not able to work out the sample type then we do not know which of the fields
            # are required so we do not report on the empty cells of the row.
            not_applicable = FIELDS_NOT_APPLICABLE.get(sample_type, [])
            cell_classes = {col: classify_cell(row[col]) for col in METADATA_FIELDS}
            if sample_type is not None:
                for col in METADATA_FIELDS:
                    if cell_classes[col] == VALUE:
                        continue
                    if col not in not_applicable:
                        self.issues.error(MSG_REQUIRED_EMPTY, cell=cell(col), note=sample_type, column=col)
                    elif cell_classes[col] == STAND_IN:
                        self.issues.error(MSG_LEAVE_BLANK, cell=cell(col), value=row[col], note=sample_type, column=col)

            # The cells that hold a value
            for col, max_degrees, range_message in [(LAT, 90, MSG_LAT_RANGE), (LON, 180, MSG_LON_RANGE)]:
                if cell_classes[col] == VALUE:
                    self._check_decimal_degrees(row[col], cell(col), col, max_degrees, range_message)
            if (cell_classes[LAT] == VALUE) != (cell_classes[LON] == VALUE):
                # If the empty one of the two is a required field then it has already been reported above
                if sample_type is not None and LAT in not_applicable and LON in not_applicable:
                    self.issues.error(MSG_LAT_LON_BOTH, cell=f'{cell(LAT)}-{cell(LON)}')

            if cell_classes['collection_date'] == VALUE:
                if EXCEL_DATE_RE.match(row['collection_date'].strip()):
                    self.issues.error(MSG_DATE_EXCEL, cell=cell('collection_date'))
                else:
                    # NB every format that we accept is also accepted by the framework
                    date_parts = accepted_date_parts(row['collection_date'])
                    if date_parts is None:
                        self.issues.error(MSG_DATE_FORMAT, cell=cell('collection_date'), value=row['collection_date'])
                    elif not date_is_real(*date_parts):
                        self.issues.error(
                            MSG_DATE_NOT_VALID, cell=cell('collection_date'), value=row['collection_date'])

            if cell_classes['collection_depth'] == VALUE:
                if EXCEL_DATE_RE.match(row['collection_depth'].strip()):
                    self.issues.error(MSG_DEPTH_EXCEL, cell=cell('collection_depth'))
                elif not DEPTH_RE.match(row['collection_depth'].strip()):
                    self.issues.error(MSG_DEPTH_FORMAT, cell=cell('collection_depth'), value=row['collection_depth'])

            if cell_classes['host_species'] == VALUE:
                # People were putting the full binomial in the speices colums.
                # The framework crops this back to just the species component, but only if there are
                # two components and the first matches the genus.
                components = row['host_species'].split(' ')
                if len(components) == 2 and components[0] == row['host_genus']:
                    self.issues.warning(
                        MSG_SPECIES_SHORTENED, cell=cell('host_species'),
                        value=row['host_species'], new_value=components[1])

            for col, max_length in MAX_FIELD_LENGTHS.items():
                if classify_cell(row[col]) == VALUE and len(row[col]) > max_length:
                    self.issues.error(f'Entry too long (max {max_length} characters)', cell=cell(col), column=col)

        for sample_name, cells in cells_of_sample_name.items():
            if len(cells) > 1:
                for _ in cells:
                    self.issues.error(MSG_NAME_DUPLICATE, cell=_, value=sample_name)
        for file_name, cells in cells_of_file_name.items():
            if len(cells) > 1:
                for _ in cells:
                    self.issues.error(MSG_FILE_DUPLICATE, cell=_, value=file_name)

        return sample_names

    def _check_sample_name(self, raw_name, cell):
        """Returns the sample name as it will be used by the framework or None if there is no valid name."""
        cell_class = classify_cell(raw_name)
        if cell_class != VALUE:
            # NB the framework drops the row if it reads the sample name as null
            if cell_class == BLANK:
                self.issues.error(MSG_NAME_MISSING, cell=cell)
            else:
                self.issues.error(MSG_NAME_MISSING, cell=cell, value=raw_name)
            return None
        stripped_name = raw_name.strip()
        # The framework removes white space and greek characters from the names
        sample_name = stripped_name.replace(' ', '_').replace('/', '_').replace('α', 'alpha').replace('β', 'beta')
        if sample_name != stripped_name:
            self.issues.warning(MSG_NAME_CHANGED, cell=cell, value=stripped_name, new_value=sample_name)
        if not set(sample_name).issubset(VALID_SAMPLE_NAME_CHARS):
            self.issues.error(MSG_NAME_CHARS, cell=cell, value=stripped_name)
        return sample_name

    def _check_file_name(self, raw_name, cell, missing_message):
        """Returns the file name or None if there is no valid name.
        NB Users are not able to select fastq or fq files for upload so the files must be compressed."""
        if classify_cell(raw_name) != VALUE:
            self.issues.error(missing_message, cell=cell)
            return None
        file_name = raw_name.strip()
        if file_name != ntpath.basename(file_name) or '/' in file_name:
            self.issues.error(MSG_FILE_PATH, cell=cell, value=file_name)
            return None
        if not file_name.endswith('fastq.gz') and not file_name.endswith('fq.gz'):
            self.issues.error(MSG_FILE_ENDING, cell=cell, value=file_name)
            return None
        return file_name

    def _check_sample_type(self, raw_type, cell):
        """Returns the sample type as it is written in the template or None if it is not one of the
        accepted types."""
        if classify_cell(raw_type) != VALUE:
            self.issues.error(MSG_REQUIRED_EMPTY, cell=cell, column='sample_type')
            return None
        for sample_type in SAMPLE_TYPES:
            if raw_type.strip().lower() == sample_type.lower():
                if raw_type.strip() != sample_type:
                    self.issues.warning(MSG_TYPE_CASE, cell=cell, value=raw_type.strip(), new_value=sample_type)
                return sample_type
        self.issues.error(MSG_TYPE_NOT_ACCEPTED, cell=cell, value=raw_type)
        return None

    def _check_decimal_degrees(self, value, cell, col, max_degrees, range_message):
        """The framework does not stop on a bad lat or lon value. It sets both to 999.
        We want the user to correct the value instead."""
        value = value.strip()
        if not DECIMAL_DEGREES_RE.match(value):
            self.issues.error(MSG_LAT_LON_NOT_DECIMAL, cell=cell, value=value, column=col)
        elif abs(float(value)) > max_degrees:
            self.issues.error(range_message, cell=cell, value=value)

    def _check_seq_files_exist(self):
        """
        Check that all of the sequencing files provided in the datasheet exist.
        Check that there are no additional files staged that are not in the datasheet
        Check that no files have been staged more than once
        Also check for size of file and require a 300B minimum. Remove from the sample from the data sheet if
        smaller than this.
        """

        # Get a list of the files that the user has added
        data_dict = json.loads(list(self.request.form.keys())[0])
        # A dict of filename to size
        self.added_files_dict = {}
        for file_dict in data_dict["files"]:
            for k, v in file_dict.items():
                self.duplicate_staged_files[k] += 1
                self.added_files_dict[k] = v

        # Delete any entries from the duplicate samples dict that are 1 and then remove
        for k_1 in [k for k, v in self.duplicate_staged_files.items() if v == 1]:
            del self.duplicate_staged_files[k_1]
        # Convert to plain dict
        # This dict will now be empty if there were no duplicate files.
        self.duplicate_staged_files = dict(self.duplicate_staged_files)

        self.file_not_found_list = []
        self.size_violation_samples = []
        checked_files = []
        for df_ind in self.sample_meta_info_df.index.values.tolist():
            fwd_file = self.sample_meta_info_df.at[df_ind, 'fastq_fwd_file_name']
            rev_file = self.sample_meta_info_df.at[df_ind, 'fastq_rev_file_name']

            checked_files.extend([fwd_file, rev_file])

            # Check for the fwd read
            # Unlike in the framework, we will not allow .fastq files. Only .fastq.gz
            self._check_file_in_added_list(df_ind, fwd_file)
            self._check_file_in_added_list(df_ind, rev_file)

            # NB if we were unable to find either the fwd or rev read then we will not be able
            # to continue with our checks
            if df_ind in self.missing_files:
                continue

            # Check file size
            if self.added_files_dict[fwd_file] < 300 or self.added_files_dict[rev_file] < 300:
                logger.warning(f'At least one of the seq files for sample {df_ind} is less than 300 bytes in size')
                self.size_violation_samples.append(df_ind)

        # drop the rows that had size violations
        # Not really necessary in terms of checking whether files exist
        # But good to undertake this action now to ensure that it doesn't create
        # Any issues during the framework's handling of the datasheet and files
        self.sample_meta_info_df.drop(index=self.size_violation_samples, inplace=True)

        # Check for files that have been added but haven't been listed in the df
        self.extra_files = [filename for filename in self.added_files_dict.keys() if filename not in checked_files]

        if self.missing_files or self.size_violation_samples or self.extra_files or self.duplicate_staged_files:
            # Then we either have missing files, size violation files or extra files
            message = []
            if self.missing_files:
                message.append("Some of the files listed in your datasheet cannot be found.")
            if self.size_violation_samples:
                message.append("Some of your files are too small.")
            if self.extra_files:
                message.append("There are files that are not listed in your datasheet.")
            if self.duplicate_staged_files:
                message.append("Some files have been staged twice.")
            message = '\n'.join(message)
            raise AddedFilesError(
                message=message,
                data={
                    "missing_files":self.missing_files,
                    "size_violation_samples":self.size_violation_samples,
                    "extra_files":self.extra_files,
                    "duplicate_staged_files": self.duplicate_staged_files
                })
        else:
            # Otherwise there were no errors and we can proceed to check for warnings
            return

    def _check_file_in_added_list(self, df_ind, filename):
        if not filename in self.added_files_dict:
            self.file_not_found_list.append(filename)
            self.missing_files[df_ind].append(filename)

    def _get_and_save_datasheet(self):
        # Then we need to get the datasheet from the uploaded files
        dsheet_files = [k for k in self.request.files if
                        self.request.files[k].filename.endswith('.xlsx') or
                        self.request.files[k].filename.endswith('.csv')]
        assert (len(dsheet_files) == 1)
        data_sheet_key = dsheet_files[0]
        # Save the datasheet to the temp directory
        # And then read in as a dataframe to work with using the current SymPortal code
        file = self.request.files.get(data_sheet_key)
        filename = secure_filename(file.filename)
        self.datasheet_path = os.path.join(self.user_upload_directory, filename)
        file.save(self.datasheet_path)
        return filename
